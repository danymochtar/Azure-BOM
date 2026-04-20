"""Normalize ANY uploaded file into Claude message content blocks.

Supported kinds:
- spreadsheet: .xlsx, .xls, .csv           → text preview block + sheets dict
- pdf:        .pdf                         → `document` content block (base64)
- image:      .png, .jpg, .jpeg, .gif, .webp → `image` content block (base64)
- docx:       .docx                        → extracted text block (python-docx)
- text:       .txt, .md, .json, .yml, .yaml, .log → plain text block
- unknown: raises ValueError

Callers (classifier, extractor) treat the returned `content_blocks` as the
leading segments of the user message, then append their own instruction
text after.

For spreadsheets the existing preview helper from ai_inventory is reused so
the token-saving direct/mapping split for large files keeps working.
"""
from __future__ import annotations

import base64
import io
import os
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

import pandas as pd


SPREADSHEET_EXTS = {".xlsx", ".xls", ".csv"}
PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
DOCX_EXTS = {".docx"}
TEXT_EXTS = {".txt", ".md", ".json", ".yml", ".yaml", ".log"}

IMAGE_MIME: Dict[str, str] = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

# Practical cap to protect both token budget and HTTP payload size.
MAX_FILE_MB = 30


@dataclass
class UploadContent:
    kind: str                      # spreadsheet | pdf | image | docx | text
    filename: str
    content_blocks: List[dict] = field(default_factory=list)
    sheets: Optional[Dict[str, pd.DataFrame]] = None  # spreadsheet only
    text_summary: str = ""
    row_count: int = 0             # spreadsheet only


def kind_from_name(filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext in SPREADSHEET_EXTS:
        return "spreadsheet"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in DOCX_EXTS:
        return "docx"
    if ext in TEXT_EXTS:
        return "text"
    return "unknown"


def _image_mime(filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    return IMAGE_MIME.get(ext, "image/png")


def _extract_docx_text(data: bytes) -> str:
    from docx import Document  # type: ignore

    doc = Document(io.BytesIO(data))
    parts: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Include tables too — design docs often put specs in tables.
    for table in doc.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))
    return "\n".join(parts)


def prepare(
    data: bytes,
    filename: str,
    spreadsheet_preview_rows: Optional[int] = None,
) -> UploadContent:
    """Build Claude-ready content blocks for any supported file type.

    For spreadsheets, `spreadsheet_preview_rows=None` sends every row (direct
    extraction) and an integer sends just a sample (mapping / classifier).
    """
    size_mb = len(data) / (1024 * 1024)
    if size_mb > MAX_FILE_MB:
        raise ValueError(
            f"File too large: {size_mb:.1f} MB (cap {MAX_FILE_MB} MB)."
        )

    kind = kind_from_name(filename)
    uc = UploadContent(kind=kind, filename=filename)

    if kind == "spreadsheet":
        # Defer to the existing helpers so any preview tweaks stay in one place.
        from .ai_inventory import _build_preview, _read_file

        sheets = _read_file(data, filename)
        uc.sheets = sheets
        uc.row_count = sum(len(df) for df in sheets.values())
        preview = _build_preview(sheets, sample_rows=spreadsheet_preview_rows)
        uc.content_blocks = [{"type": "text", "text": preview}]
        uc.text_summary = (
            f"Spreadsheet — {len(sheets)} sheet(s), {uc.row_count:,} rows total"
        )
        return uc

    if kind == "pdf":
        b64 = base64.standard_b64encode(data).decode()
        uc.content_blocks = [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": b64,
                },
            }
        ]
        uc.text_summary = f"PDF — {len(data):,} bytes"
        return uc

    if kind == "image":
        b64 = base64.standard_b64encode(data).decode()
        uc.content_blocks = [
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": _image_mime(filename),
                    "data": b64,
                },
            }
        ]
        uc.text_summary = f"Image — {len(data):,} bytes ({_image_mime(filename)})"
        return uc

    if kind == "docx":
        try:
            text = _extract_docx_text(data)
        except Exception as e:
            raise ValueError(f"Failed to read DOCX: {e}") from e
        uc.content_blocks = [
            {"type": "text", "text": f"<docx_content>\n{text}\n</docx_content>"}
        ]
        uc.text_summary = f"DOCX — {len(text):,} chars extracted"
        return uc

    if kind == "text":
        text = data.decode("utf-8", errors="replace")
        uc.content_blocks = [
            {"type": "text", "text": f"<file_content>\n{text}\n</file_content>"}
        ]
        uc.text_summary = f"Text — {len(text):,} chars"
        return uc

    raise ValueError(
        f"Unsupported file type: {filename}. Supported: "
        f"{sorted(SPREADSHEET_EXTS | PDF_EXTS | IMAGE_EXTS | DOCX_EXTS | TEXT_EXTS)}"
    )


ALL_SUPPORTED_EXTS: List[str] = sorted(
    e.lstrip(".")
    for e in (SPREADSHEET_EXTS | PDF_EXTS | IMAGE_EXTS | DOCX_EXTS | TEXT_EXTS)
)
